from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Count
from django.core.mail import send_mail
from django.conf import settings
from django.core.exceptions import ObjectDoesNotExist
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.permissions import AllowAny
from django.utils import timezone
import openpyxl
import uuid
import random
import string
from django.contrib.auth.hashers import make_password

from .models import (
    Usuario, Persona, AreaTrabajo, Carrera, Institucion,
    Empresa, ProgramaFormacion, Aspirante, Vacante, Postulacion, Curriculo,
    Practicante, Notificacion, Auditoria, Capacitacion,
    Favorito, InscripcionCapacitacion, Entrevista
)

from .serializers import (
    UsuarioSerializer, PersonaSerializer, AreaTrabajoSerializer,
    CarreraSerializer, InstitucionSerializer, EmpresaSerializer,
    ProgramaFormacionSerializer, AspiranteSerializer, VacanteSerializer,
    PostulacionSerializer, CurriculoSerializer, PracticanteSerializer,
    NotificacionSerializer, AuditoriaSerializer, CapacitacionSerializer,
    CrearPerfilAspiranteSerializer, CrearPerfilEmpresaSerializer, CrearPerfilInstitucionSerializer, LoginSerializer, MiPerfilSerializer,
    FavoritoSerializer, InscripcionCapacitacionSerializer, EntrevistaSerializer
)

from django.views.generic import TemplateView

class FrontendAppView(TemplateView):
    template_name = "index.html"

# =====================================================
# ENVIAR CREDENCIALES AL ASPIRANTE (EMAIL)
# =====================================================

class EnviarCredencialesAspiranteView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        correo = request.data.get('correo')
        contrasena = request.data.get('contrasena')
        nombre = request.data.get('nombre', 'Estudiante')
        institucion = request.data.get('institucion', 'tu institución')

        if not correo or not contrasena:
            return Response({'error': 'Correo y contraseña son requeridos'}, status=status.HTTP_400_BAD_REQUEST)

        asunto = '🎉 ¡Bienvenido a GreenTalent! Tus credenciales de acceso'

        cuerpo_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
          <meta charset="UTF-8">
          <style>
            body {{ font-family: Arial, sans-serif; background: #f4f4f4; margin: 0; padding: 20px; }}
            .container {{ max-width: 600px; margin: 0 auto; background: white; border-radius: 16px; overflow: hidden; box-shadow: 0 4px 20px rgba(0,0,0,0.1); }}
            .header {{ background: linear-gradient(135deg, #1a8641, #0f5e2f); padding: 40px 30px; text-align: center; }}
            .header h1 {{ color: white; margin: 0; font-size: 28px; font-weight: 900; letter-spacing: -1px; }}
            .header p {{ color: rgba(255,255,255,0.8); margin: 8px 0 0; font-size: 14px; }}
            .body {{ padding: 40px 30px; }}
            .greeting {{ font-size: 18px; font-weight: bold; color: #1a2236; margin-bottom: 12px; }}
            .text {{ color: #555; line-height: 1.6; font-size: 14px; margin-bottom: 20px; }}
            .credentials-box {{ background: #f8fffe; border: 2px solid #1a8641; border-radius: 12px; padding: 24px; margin: 24px 0; }}
            .credentials-box .label {{ font-size: 11px; font-weight: 900; color: #1a8641; text-transform: uppercase; letter-spacing: 1px; margin-bottom: 4px; }}
            .credentials-box .value {{ font-size: 16px; font-weight: bold; color: #1a2236; background: white; border: 1px solid #e0e0e0; padding: 10px 16px; border-radius: 8px; margin-bottom: 16px; font-family: monospace; }}
            .btn {{ display: inline-block; background: #1a8641; color: white; text-decoration: none; padding: 14px 32px; border-radius: 10px; font-weight: 900; font-size: 14px; text-transform: uppercase; letter-spacing: 1px; }}
            .footer {{ padding: 20px 30px; text-align: center; background: #f9f9f9; color: #aaa; font-size: 11px; }}
          </style>
        </head>
        <body>
          <div class="container">
            <div class="header">
              <h1>🌿 GreenTalent</h1>
              <p>Zona Franca La Lima · Plataforma de Talento</p>
            </div>
            <div class="body">
              <p class="greeting">¡Hola, {nombre}! 👋</p>
              <p class="text">
                Has sido registrado(a) en <strong>GreenTalent</strong> por <strong>{institucion}</strong>.
                A continuación encontrarás tus credenciales de acceso a la plataforma.
              </p>

              <div class="credentials-box">
                <div class="label">Correo de acceso</div>
                <div class="value">{correo}</div>
                <div class="label">Contraseña temporal</div>
                <div class="value">{contrasena}</div>
              </div>

              <p class="text">
                Te recomendamos ingresar y <strong>cambiar tu contraseña</strong> lo antes posible desde la sección de Ajustes.
              </p>

              <a href="http://localhost:5173/login" class="btn">Iniciar sesión ahora →</a>

              <p class="text" style="margin-top:24px; font-size:12px; color:#999;">
                Si no esperabas este correo, puedes ignorarlo. Si tienes dudas, contacta a tu institución.
              </p>
            </div>
            <div class="footer">
              © 2025 GreenTalent · Zona Franca La Lima · Todos los derechos reservados
            </div>
          </div>
        </body>
        </html>
        """

        cuerpo_texto = f"""
        ¡Bienvenido a GreenTalent, {nombre}!

        Has sido registrado por: {institucion}

        Tus credenciales:
        - Correo: {correo}
        - Contraseña temporal: {contrasena}

        Ingresa en: http://localhost:5173/login

        Recuerda cambiar tu contraseña después de iniciar sesión.
        """

        try:
            send_mail(
                subject=asunto,
                message=cuerpo_texto,
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[correo],
                html_message=cuerpo_html,
                fail_silently=False,
            )
            return Response({'mensaje': f'Correo enviado a {correo}'}, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


# =====================================================
# MI PERFIL (GET)
# =====================================================

class MiPerfilView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, usuario_id=None):
        uid = usuario_id or request.query_params.get('usuario_id')
        if not uid:
            return Response({"error": "Falta usuario_id"}, status=status.HTTP_400_BAD_REQUEST)
            
        try:
            usuario = Usuario.objects.get(id=uid)
        except (Usuario.DoesNotExist, ValueError):
            # Si el ID no existe pero tal vez es un usuario eliminado del navegador, 
            # podríamos intentar buscar otro usuario con el mismo ID en el futuro o manejarlo mejor.
            return Response(
                {"error": "Usuario no existe"},
                status=status.HTTP_404_NOT_FOUND
            )

        if usuario.rol == 'empresa':
            empresa = usuario.empresa_set.first()
            if not empresa:
                # Recuperar de otro usuario con el mismo correo
                otro_u = Usuario.objects.filter(correo__iexact=usuario.correo, rol='empresa').exclude(id=usuario.id).first()
                if otro_u:
                    empresa = otro_u.empresa_set.first()
            
            if not empresa:
                return Response({"usuario_id": str(usuario.id), "perfil_completo": False, "rol": "empresa"}, status=status.HTTP_200_OK)
            
            return Response({
                "usuario_id": str(usuario.id),
                "empresa_id": str(empresa.id),
                "nombre": empresa.nombre,
                "descripcion": empresa.descripcion,
                "nombre_contacto": empresa.nombre_contacto,
                "correo_contacto": empresa.correo_contacto,
                "url_externa": empresa.url_externa,
                "ubicacion": empresa.ubicacion,
                "sector": empresa.sector,
                "tamano_empresa": empresa.tamano_empresa,
                "logo_url": request.build_absolute_uri(empresa.logo_url.url) if empresa.logo_url else None,
                "perfil_completo": True,
                "rol": "empresa"
            }, status=status.HTTP_200_OK)

        if usuario.rol == 'institucion':
            institucion = usuario.institucion_set.first()
            if not institucion:
                # Recuperar de otro usuario con el mismo correo
                otro_u = Usuario.objects.filter(correo__iexact=usuario.correo, rol='institucion').exclude(id=usuario.id).first()
                if otro_u:
                    institucion = otro_u.institucion_set.first()

            if not institucion:
                return Response({"usuario_id": str(usuario.id), "perfil_completo": False, "rol": "institucion"}, status=status.HTTP_200_OK)

            # Practicantes vinculados a esta institución
            practicantes = Practicante.objects.filter(institucion=institucion).select_related('aspirante__usuario__persona')
            practicantes_data = []
            for p in practicantes:
                estudiante = p.aspirante
                try:
                    persona = estudiante.usuario.persona
                    nom = f"{persona.nombre} {persona.apellidos}"
                    correo = estudiante.usuario.correo
                except:
                    nom = "Estudiante"
                    correo = ""
                practicantes_data.append({
                    "id": str(p.id),
                    "nombre": nom,
                    "correo": correo,
                    "programa": p.nombre_programa,
                    "nivel_academico": p.nivel_academico,
                    "estado": p.estado_pasantia,
                    "horas_requeridas": p.horas_requeridas,
                    "fecha_inicio": p.fecha_inicio.isoformat() if p.fecha_inicio else None,
                    "fecha_fin": p.fecha_fin.isoformat() if p.fecha_fin else None,
                })

            return Response({
                "usuario_id": str(usuario.id),
                "institucion_id": str(institucion.id),
                "nombre": institucion.nombre,
                "titulo": institucion.titulo,
                "tipo": institucion.tipo,
                "nombre_contacto": institucion.nombre_contacto,
                "correo_contacto": institucion.correo_contacto,
                "correo": usuario.correo,
                "telefono": usuario.telefono,
                "activa": institucion.activa,
                "total_practicantes": len(practicantes_data),
                "practicantes": practicantes_data,
                "perfil_completo": True,
                "rol": "institucion"
            }, status=status.HTTP_200_OK)

        if usuario.rol == 'admin':
            return Response({
                "usuario_id": str(usuario.id),
                "nombre": "Administrador",
                "apellidos": "Sistema",
                "correo": usuario.correo,
                "rol": "admin",
                "perfil_completo": True,
                "preferencias": usuario.preferencias
            }, status=status.HTTP_200_OK)

        # Si el rol es aspirante, buscamos Persona + Aspirante
        if usuario.rol == 'aspirante':
            persona = getattr(usuario, 'persona', None)
            aspirante = getattr(usuario, 'aspirante', None)

            if not persona or not aspirante:
                # INTENTO DE RECUPERACIÓN: Si hay otro usuario con el mismo correo que SÍ tenga perfil
                # Buscamos una Persona que tenga Aspirante y cuyo usuario tenga el mismo correo
                otra_persona = Persona.objects.filter(usuario__correo__iexact=usuario.correo).exclude(usuario=usuario).filter(aspirante__isnull=False).first()
                
                if otra_persona:
                    persona = otra_persona
                    aspirante = otra_persona.aspirante
                    # Usamos el correo del usuario actual pero los datos de la otra persona
                else:
                    # Perfil realmente incompleto: devolver un objeto seguro para el frontend
                    return Response({
                        "usuario_id": str(usuario.id),
                        "persona_id": None,
                        "aspirante_id": None,
                        "nombre": "Usuario",
                        "apellidos": "Registrado",
                        "cedula": "",
                        "correo": usuario.correo,
                        "foto_url": None,
                        "sobre_mi": "",
                        "carrera": "Estudiante",
                        "carrera_id": None,
                        "nivel_educativo": "",
                        "estado_laboral": "",
                        "telefono": "",
                        "provincia": "",
                        "canton": "",
                        "genero": "",
                        "nacionalidad": "Costarricense",
                        "fecha_nacimiento": None,
                        "habilidades_tecnicas": [],
                        "habilidades_blandas": [],
                        "experiencia": [],
                        "practicante": False,
                        "postulaciones": [],
                        "favoritos": [],
                        "capacitaciones_inscritas": [],
                        "entrevistas": [],
                        "preferencias": usuario.preferencias,
                        "institucion_origen": None,
                        "perfil_completo": False,
                        "rol": "aspirante"
                    }, status=status.HTTP_200_OK)

            postulaciones = Postulacion.objects.filter(aspirante=aspirante).select_related('vacante', 'vacante__empresa')
            postulaciones_data = []
            for p in postulaciones:
                postulaciones_data.append({
                    "id": p.id,
                    "cargo": p.vacante.titulo,
                    "empresa": p.vacante.empresa.nombre,
                    "tipo": p.vacante.tipo_vacante,
                    "estado": p.estado,
                    "fecha": p.postulado_en.isoformat() if p.postulado_en else None
                })

            # Aplanamos la respuesta para que sea más fácil de usar en el frontend
            data = {
                "usuario_id": str(usuario.id),
                "persona_id": str(persona.id),
                "aspirante_id": str(aspirante.id),
                "nombre": persona.nombre,
                "apellidos": persona.apellidos,
                "cedula": persona.cedula,
                "correo": usuario.correo,
                "foto_url": request.build_absolute_uri(aspirante.foto_url.url) if aspirante.foto_url else None,
                "sobre_mi": aspirante.sobre_mi,
                "carrera": aspirante.carrera.nombre if aspirante.carrera else "Estudiante",
                "carrera_id": aspirante.carrera.id if aspirante.carrera else None,
                "nivel_educativo": aspirante.nivel_educativo,
                "estado_laboral": aspirante.estado_laboral,
                "telefono": persona.telefono,
                "provincia": persona.provincia,
                "canton": persona.canton,
                "genero": persona.genero,
                "nacionalidad": persona.nacionalidad,
                "fecha_nacimiento": persona.fecha_nacimiento.isoformat() if persona.fecha_nacimiento else None,
                "habilidades_tecnicas": aspirante.habilidades_tecnicas,
                "habilidades_blandas": aspirante.habilidades_blandas,
                "experiencia": aspirante.experiencia,
                "practicante": hasattr(aspirante, 'practicante_set') and aspirante.practicante_set.exists(),
                "postulaciones": postulaciones_data,
                "favoritos": [],
                "capacitaciones_inscritas": [],
                "entrevistas": [],
                "preferencias": usuario.preferencias,
                "password_plano": usuario.password_plano,
                "institucion_origen": {
                    "id": str(aspirante.institucion_origen.id),
                    "nombre": aspirante.institucion_origen.nombre
                } if aspirante.institucion_origen else None,
                "perfil_completo": True,
                "rol": "aspirante"
            }

            # --- FAVORITOS ---
            favoritos = Favorito.objects.filter(aspirante=aspirante).select_related(
                'vacante', 'capacitacion', 'vacante__empresa', 'capacitacion__empresa', 'capacitacion__institucion'
            )
            for fav in favoritos:
                entidad = "N/A"
                if fav.vacante and fav.vacante.empresa:
                    entidad = fav.vacante.empresa.nombre
                elif fav.capacitacion:
                    entidad = fav.capacitacion.empresa.nombre if fav.capacitacion.empresa else (fav.capacitacion.institucion.nombre if fav.capacitacion.institucion else "N/A")
                
                data["favoritos"].append({
                    "id": fav.id,
                    "tipo": "vacante" if fav.vacante else "capacitacion",
                    "obj_id": fav.vacante.id if fav.vacante else fav.capacitacion.id,
                    "titulo": fav.vacante.titulo if fav.vacante else fav.capacitacion.titulo,
                    "entidad": entidad
                })

            # --- CAPACITACIONES INSCRITAS ---
            inscripciones = InscripcionCapacitacion.objects.filter(aspirante=aspirante).select_related(
                'capacitacion', 'capacitacion__empresa', 'capacitacion__institucion'
            )
            for ins in inscripciones:
                entidad = ins.capacitacion.empresa.nombre if ins.capacitacion.empresa else (ins.capacitacion.institucion.nombre if ins.capacitacion.institucion else "N/A")
                data["capacitaciones_inscritas"].append({
                    "id": ins.id,
                    "cap_id": ins.capacitacion.id,
                    "titulo": ins.capacitacion.titulo,
                    "entidad": entidad,
                    "fecha_inscripcion": ins.fecha_inscripcion.isoformat()
                })

            # --- INSTITUCIÓN DE ORIGEN (via Practicante) ---
            practicante_obj = aspirante.practicante_set.select_related('institucion').first()
            if practicante_obj and practicante_obj.institucion:
                inst = practicante_obj.institucion
                data["institucion_origen"] = {
                    "id": str(inst.id),
                    "nombre": inst.nombre,
                    "titulo": inst.titulo,
                    "tipo": inst.tipo,
                    "nombre_contacto": inst.nombre_contacto,
                    "correo_contacto": inst.correo_contacto,
                    "programa": practicante_obj.nombre_programa,
                    "nivel_academico": practicante_obj.nivel_academico,
                    "horas_requeridas": practicante_obj.horas_requeridas,
                    "estado_pasantia": practicante_obj.estado_pasantia,
                    "fecha_inicio": practicante_obj.fecha_inicio.isoformat() if practicante_obj.fecha_inicio else None,
                    "fecha_fin": practicante_obj.fecha_fin.isoformat() if practicante_obj.fecha_fin else None,
                }

            # --- ENTREVISTAS ---
            entrevistas = Entrevista.objects.filter(aspirante=aspirante).select_related('postulacion__vacante', 'empresa')
            for ent in entrevistas:
                data["entrevistas"].append({
                    "id": str(ent.id),
                    "cargo": ent.postulacion.vacante.titulo if ent.postulacion and ent.postulacion.vacante else "N/A",
                    "empresa": ent.empresa.nombre if ent.empresa else "Empresa",
                    "fecha": ent.fecha.isoformat(),
                    "hora": ent.hora.strftime("%H:%M") if ent.hora else "N/A",
                    "meet_url": ent.meet_url,
                    "estado": ent.estado
                })

            return Response(data, status=status.HTTP_200_OK)

        return Response({"error": "Rol no reconocido"}, status=status.HTTP_400_BAD_REQUEST)


# =====================================================
# CREAR PERFIL ASPIRANTE
# =====================================================

class CrearPerfilAspiranteView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = CrearPerfilAspiranteSerializer(data=request.data)

        if serializer.is_valid():
            aspirante = serializer.save()
            return Response({
                "mensaje": "Perfil creado correctamente",
                "aspirante_id": aspirante.id
            }, status=status.HTTP_201_CREATED)

        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

# =====================================================
# ENVIAR CREDENCIALES ASPIRANTE
# =====================================================

from django.core.mail import send_mail

class EnviarCredencialesAspiranteView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        correo = request.data.get('correo')
        contrasena = request.data.get('contrasena')
        nombre = request.data.get('nombre', 'Estudiante')
        institucion = request.data.get('institucion', 'una institución')

        if not correo or not contrasena:
            return Response({'error': 'Faltan datos requeridos.'}, status=status.HTTP_400_BAD_REQUEST)

        asunto = f"¡Bienvenido a GreenTalent, {nombre}!"
        mensaje = f"""Hola {nombre},

Has sido recomendado en la plataforma de GreenTalent por {institucion}.
Nos alegra contar con tu talento en nuestro ecosistema.

Tus credenciales de acceso temporal son:
Usuario/Correo: {correo}
Contraseña: {contrasena}

Por favor, inicia sesión lo antes posible y completa tu perfil profesional.

Atentamente,
El equipo de GreenTalent"""
        
        try:
            send_mail(asunto, mensaje, 'noreply@greentalent.com', [correo], fail_silently=False)
            return Response({'status': 'Correo enviado exitosamente'}, status=status.HTTP_200_OK)
        except Exception as e:
            print(f"Error enviando correo: {e}")
            return Response({'status': 'Registro completado, pero sin envío de correo'}, status=status.HTTP_200_OK)


# =====================================================
# CREAR PERFIL EMPRESA
# =====================================================

class CrearPerfilEmpresaView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = CrearPerfilEmpresaSerializer(data=request.data)
        if serializer.is_valid():
            empresa = serializer.save()
            return Response({
                "mensaje": "Perfil de empresa creado correctamente",
                "empresa_id": empresa.id
            }, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class CrearPerfilInstitucionView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = CrearPerfilInstitucionSerializer(data=request.data)
        if serializer.is_valid():
            inst = serializer.save()
            return Response({"mensaje": "Perfil de institución creado", "id": inst.id}, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

# =====================================================
# CONFIGURACIÓN Y AJUSTES
# =====================================================

class UpdatePreferenciasView(APIView):
    permission_classes = [AllowAny]

    def post(self, request, usuario_id):
        try:
            usuario = Usuario.objects.get(id=usuario_id)
            usuario.preferencias = request.data.get('preferencias', {})
            usuario.save()
            return Response({"mensaje": "Preferencias actualizadas"}, status=status.HTTP_200_OK)
        except Usuario.DoesNotExist:
            return Response({"error": "Usuario no encontrado"}, status=status.HTTP_404_NOT_FOUND)

class CambiarPasswordView(APIView):
    permission_classes = [AllowAny]

    def post(self, request, usuario_id):
        from django.contrib.auth.hashers import make_password, check_password
        try:
            usuario = Usuario.objects.get(id=usuario_id)
            current_password = request.data.get('current_password')
            new_password = request.data.get('new_password')

            if not check_password(current_password, usuario.contrasena_hash):
                return Response({"error": "La contraseña actual es incorrecta"}, status=status.HTTP_400_BAD_REQUEST)

            usuario.contrasena_hash = make_password(new_password)
            usuario.password_plano = None
            usuario.save()
            return Response({"mensaje": "Contraseña actualizada correctamente"}, status=status.HTTP_200_OK)
        except Usuario.DoesNotExist:
            return Response({"error": "Usuario no encontrado"}, status=status.HTTP_404_NOT_FOUND)

class EliminarCuentaView(APIView):
    permission_classes = [AllowAny]

    def delete(self, request, usuario_id):
        try:
            usuario = Usuario.objects.get(id=usuario_id)
            usuario.delete()
            return Response({"mensaje": "Cuenta eliminada correctamente"}, status=status.HTTP_200_OK)
        except Usuario.DoesNotExist:
            return Response({"error": "Usuario no encontrado"}, status=status.HTTP_404_NOT_FOUND)

# =====================================================
# LOGIN
# =====================================================

class LoginView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        print("LOGIN REQUEST DATA:", request.data)
        serializer = LoginSerializer(data=request.data)
        if not serializer.is_valid():
            print("LOGIN ERRORS:", serializer.errors)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        return Response(serializer.validated_data, status=status.HTTP_200_OK)


# =====================================================
# CRUD AUTOMÁTICO CON VIEWSETS
# =====================================================

class UsuarioViewSet(viewsets.ModelViewSet):
    queryset = Usuario.objects.all()
    serializer_class = UsuarioSerializer

    @action(detail=True, methods=['post'])
    def dar_consentimiento(self, request, pk=None):
        usuario = self.get_object()
        usuario.consentimiento = True
        usuario.fecha_consentimiento = timezone.now()
        usuario.save()
        return Response({'status': 'Consentimiento registrado'})

    @action(detail=True, methods=['delete'])
    def eliminar_cuenta(self, request, pk=None):
        usuario = self.get_object()
        usuario.activo = False
        usuario.save()
        return Response({'status': 'Cuenta desactivada'})

    @action(detail=False, methods=['get'])
    def stats(self, request):
        from .models import Aspirante, Empresa, Vacante, Postulacion
        from django.db.models.functions import TruncMonth
        from django.utils import timezone
        import datetime
        
        # Totales básicos
        total_estudiantes = Aspirante.objects.count()
        total_empresas = Empresa.objects.count()
        total_vacantes = Vacante.objects.count()
        total_postulaciones = Postulacion.objects.count()
        
        # Distribución de vacantes por área
        vacantes_por_area = Vacante.objects.values('area_trabajo__nombre').annotate(
            count=Count('id')
        ).order_by('-count')
        
        # Distribución de estudiantes por carrera
        estudiantes_por_carrera = Aspirante.objects.values('carrera__nombre').annotate(
            count=Count('id')
        ).order_by('-count')
        
        # Tasa de colocación (simulada o real si tenemos el campo)
        colocados = Aspirante.objects.filter(estado_laboral='empleado').count()
        tasa_colocacion = (colocados / total_estudiantes * 100) if total_estudiantes > 0 else 0

        # Tendencias de registro (últimos 6 meses)
        seis_meses_atras = timezone.now() - datetime.timedelta(days=180)
        
        est_trends = Aspirante.objects.filter(creado_en__gte=seis_meses_atras) \
            .annotate(month=TruncMonth('creado_en')) \
            .values('month') \
            .annotate(count=Count('id')) \
            .order_by('month')

        emp_trends = Empresa.objects.filter(creado_en__gte=seis_meses_atras) \
            .annotate(month=TruncMonth('creado_en')) \
            .values('month') \
            .annotate(count=Count('id')) \
            .order_by('month')

        # Combinar tendencias
        tendencias = []
        labels = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
        
        # Simplificación para el dashboard (solo los últimos meses encontrados)
        trends_map = {}
        for item in est_trends:
            m = item['month'].month
            label = labels[m-1]
            trends_map[label] = {"month": label, "estudiantes": item['count'], "empresas": 0}
            
        for item in emp_trends:
            m = item['month'].month
            label = labels[m-1]
            if label in trends_map:
                trends_map[label]["empresas"] = item['count']
            else:
                trends_map[label] = {"month": label, "estudiantes": 0, "empresas": item['count']}
        
        tendencias = sorted(trends_map.values(), key=lambda x: list(trends_map.keys()).index(x['month']))

        # Distribución de roles
        roles_dist = [
            {"name": "Estudiantes", "value": total_estudiantes},
            {"name": "Empresas", "value": total_empresas},
            {"name": "Instituciones", "value": Usuario.objects.filter(rol='institucion').count()}
        ]

        return Response({
            "total_estudiantes": total_estudiantes,
            "total_empresas": total_empresas,
            "total_vacantes": total_vacantes,
            "total_postulaciones": total_postulaciones,
            "vacantes_por_area": list(vacantes_por_area),
            "estudiantes_por_carrera": list(estudiantes_por_carrera),
            "tasa_colocacion": round(tasa_colocacion, 2),
            "tendencia_registros": tendencias,
            "distribucion_roles": roles_dist
        })


class PersonaViewSet(viewsets.ModelViewSet):
    queryset = Persona.objects.all()
    serializer_class = PersonaSerializer


class AreaTrabajoViewSet(viewsets.ModelViewSet):
    queryset = AreaTrabajo.objects.all()
    serializer_class = AreaTrabajoSerializer


class CarreraViewSet(viewsets.ModelViewSet):
    queryset = Carrera.objects.all()
    serializer_class = CarreraSerializer


class InstitucionViewSet(viewsets.ModelViewSet):
    queryset = Institucion.objects.all()
    serializer_class = InstitucionSerializer

    @action(detail=True, methods=['get'])
    def descargar_plantilla(self, request, pk=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Estudiantes"
        headers = ['Nombre', 'Apellidos', 'Cedula', 'Correo', 'Telefono', 'Nivel Educativo', 'Carrera ID']
        ws.append(headers)
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=plantilla_estudiantes.xlsx'
        wb.save(response)
        return response

    @action(detail=True, methods=['post'])
    def cargar_excel(self, request, pk=None):
        institucion = self.get_object()
        file = request.FILES.get('file')
        
        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            created_count = 0
            errors = []
            
            headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
            required = ['Nombre', 'Apellidos', 'Cedula', 'Correo', 'Telefono', 'Nivel Educativo']
            
            # Buscamos columnas ignorando mayúsculas/minúsculas y espacios extras
            def clean_str(s):
                return s.replace(" ", "").lower()
            
            header_map = {clean_str(h): idx for idx, h in enumerate(headers)}
            missing = [req for req in required if clean_str(req) not in header_map]
            
            if missing:
                 return Response({"error": f"Faltan columnas requeridas o tienen nombres incorrectos: {missing}"}, status=status.HTTP_400_BAD_REQUEST)

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]: continue
                try:
                    nombre, apellidos, cedula, correo, telefono, nivel_educativo, carrera_id = row[:7]
                    
                    # Generar contraseña aleatoria
                    password_temp = ''.join(random.choice(string.ascii_letters + string.digits) for i in range(8))
                    
                    usuario = Usuario.objects.create(
                        correo=correo,
                        telefono=telefono,
                        rol='aspirante',
                        contrasena_hash=make_password(password_temp),
                        password_plano=password_temp
                    )
                    
                    persona = Persona.objects.create(
                        usuario=usuario,
                        nombre=nombre,
                        apellidos=apellidos,
                        cedula=cedula,
                        telefono=telefono,
                        nacionalidad='Honduras',
                        genero='No especificado',
                        provincia='N/A',
                        canton='N/A'
                    )
                    
                    carrera = None
                    if carrera_id:
                        try:
                            carrera = Carrera.objects.get(id=carrera_id)
                        except Carrera.DoesNotExist:
                            pass
                            
                    Aspirante.objects.create(
                        usuario=usuario,
                        persona=persona,
                        carrera=carrera,
                        institucion_origen=institucion,
                        nivel_educativo=nivel_educativo,
                        estado_laboral='buscando'
                    )
                    created_count += 1
                except Exception as e:
                    errors.append(f"Error en fila {row_idx}: {str(e)}")
            
            return Response({
                "message": f"{created_count} estudiantes cargados exitosamente.",
                "errors": errors if errors else None
            })
            
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['delete'])
    def eliminar_base_datos(self, request, pk=None):
        institucion = self.get_object()
        aspirantes = Aspirante.objects.filter(institucion_origen=institucion)
        count = aspirantes.count()
        usuarios_ids = list(aspirantes.values_list('usuario_id', flat=True))
        Usuario.objects.filter(id__in=usuarios_ids).delete()
        
        return Response({"message": f"Se eliminaron {count} perfiles de estudiantes cargados por la institución."})

    @action(detail=True, methods=['get'])
    def metricas(self, request, pk=None):
        demand = Vacante.objects.values('area_trabajo__nombre').annotate(total=Count('id')).order_by('-total')
        return Response({
            "demanda_laboral": list(demand)
        })

    @action(detail=True, methods=['get'])
    def estudiantes(self, request, pk=None):
        institucion = self.get_object()
        aspirantes = Aspirante.objects.filter(institucion_origen=institucion).select_related('usuario', 'persona')
        serializer = AspiranteSerializer(aspirantes, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['get'])
    def entrevistas(self, request, pk=None):
        institucion = self.get_object()
        entrevistas = Entrevista.objects.filter(
            aspirante__institucion_origen=institucion
        ).select_related('empresa', 'aspirante', 'aspirante__persona')
        
        serializer = EntrevistaSerializer(entrevistas, many=True)
        return Response(serializer.data)


class EmpresaViewSet(viewsets.ModelViewSet):
    queryset = Empresa.objects.all()
    serializer_class = EmpresaSerializer


class ProgramaFormacionViewSet(viewsets.ModelViewSet):
    queryset = ProgramaFormacion.objects.all()
    serializer_class = ProgramaFormacionSerializer


class AspiranteViewSet(viewsets.ModelViewSet):
    queryset = Aspirante.objects.all()
    serializer_class = AspiranteSerializer

    @action(detail=True, methods=['post'])
    def marcar_colocado(self, request, pk=None):
        aspirante = self.get_object()
        aspirante.estado_laboral = 'empleado'
        aspirante.save()
        return Response({'status': 'Estudiante marcado como colocado'})


class VacanteViewSet(viewsets.ModelViewSet):
    queryset = Vacante.objects.all()
    serializer_class = VacanteSerializer


class PostulacionViewSet(viewsets.ModelViewSet):
    permission_classes = [AllowAny]
    queryset = Postulacion.objects.all()
    serializer_class = PostulacionSerializer

    def create(self, request, *args, **kwargs):
        # Allow aspirante field to be either aspirante_id or usuario_id
        aspirante_val = request.data.get('aspirante')
        if aspirante_val:
            try:
                # Si no existe como Aspirante ID, tratar como Usuario ID
                Aspirante.objects.get(id=aspirante_val)
            except Aspirante.DoesNotExist:
                try:
                    asp = Aspirante.objects.get(usuario__id=aspirante_val)
                    request.data['aspirante'] = str(asp.id)
                except Aspirante.DoesNotExist:
                    pass
        return super().create(request, *args, **kwargs)


class CurriculoViewSet(viewsets.ModelViewSet):
    queryset = Curriculo.objects.all()
    serializer_class = CurriculoSerializer


class PracticanteViewSet(viewsets.ModelViewSet):
    queryset = Practicante.objects.all()
    serializer_class = PracticanteSerializer


class NotificacionViewSet(viewsets.ModelViewSet):
    queryset = Notificacion.objects.all()
    serializer_class = NotificacionSerializer


class AuditoriaViewSet(viewsets.ModelViewSet):
    queryset = Auditoria.objects.all()
    serializer_class = AuditoriaSerializer


class CapacitacionViewSet(viewsets.ModelViewSet):
    permission_classes = [AllowAny]
    queryset = Capacitacion.objects.all().order_by('-creado_en')
    serializer_class = CapacitacionSerializer


class FavoritoViewSet(viewsets.ModelViewSet):
    permission_classes = [AllowAny]
    queryset = Favorito.objects.all()
    serializer_class = FavoritoSerializer

    @action(detail=False, methods=['post'])
    def toggle(self, request):
        aspirante_id = request.data.get('aspirante_id')
        vacante_id = request.data.get('vacante_id')
        capacitacion_id = request.data.get('capacitacion_id')

        if not aspirante_id:
            return Response({"error": "Falta aspirante_id"}, status=400)

        # Acepta tanto aspirante_id real como usuario_id
        try:
            aspirante_obj = Aspirante.objects.get(id=aspirante_id)
        except (Aspirante.DoesNotExist, Exception):
            try:
                # Tratar como usuario_id
                aspirante_obj = Aspirante.objects.get(usuario__id=aspirante_id)
            except Aspirante.DoesNotExist:
                return Response({"error": "Aspirante no encontrado"}, status=404)

        filtros = {'aspirante': aspirante_obj}
        if vacante_id: filtros['vacante_id'] = vacante_id
        elif capacitacion_id: filtros['capacitacion_id'] = capacitacion_id
        else: return Response({"error": "Falta vacante_id o capacitacion_id"}, status=400)

        fav_exists = Favorito.objects.filter(**filtros).first()
        if fav_exists:
            fav_exists.delete()
            return Response({"status": "removed"}, status=200)
        else:
            Favorito.objects.create(**filtros)
            return Response({"status": "added"}, status=201)


class InscripcionCapacitacionViewSet(viewsets.ModelViewSet):
    permission_classes = [AllowAny]
    queryset = InscripcionCapacitacion.objects.all()
    serializer_class = InscripcionCapacitacionSerializer

    def create(self, request, *args, **kwargs):
        # Allow aspirante field to be either aspirante_id or usuario_id
        aspirante_val = request.data.get('aspirante')
        if aspirante_val:
            try:
                # Si no existe como Aspirante ID, tratar como Usuario ID
                Aspirante.objects.get(id=aspirante_val)
            except Aspirante.DoesNotExist:
                try:
                    asp = Aspirante.objects.get(usuario__id=aspirante_val)
                    request.data['aspirante'] = str(asp.id)
                except Aspirante.DoesNotExist:
                    pass
        return super().create(request, *args, **kwargs)


class EntrevistaViewSet(viewsets.ModelViewSet):
    permission_classes = [AllowAny]
    queryset = Entrevista.objects.all()
    serializer_class = EntrevistaSerializer

    def get_queryset(self):
        queryset = Entrevista.objects.all()
        usuario_id = self.request.query_params.get('usuario_id')
        rol = self.request.query_params.get('rol')
        
        if usuario_id and rol:
            try:
                import uuid
                # Validar que sea un UUID válido antes de buscar
                uid_obj = uuid.UUID(usuario_id)
                usr = Usuario.objects.get(id=uid_obj)
                email = usr.correo
                
                if rol == 'empresa':
                    queryset = queryset.filter(empresa__usuario__correo__iexact=email)
                elif rol == 'aspirante':
                    queryset = queryset.filter(aspirante__usuario__correo__iexact=email)
            except (Usuario.DoesNotExist, ValueError):
                # Fallback: intentar filtrar por ID de todos modos por si es un ID legacy que aún existe
                # (aunque ya pusimos validación arriba, el except ValueError capturará ids mal formateados)
                if rol == 'empresa':
                    queryset = queryset.filter(empresa__usuario__id=usuario_id)
                elif rol == 'aspirante':
                    queryset = queryset.filter(aspirante__usuario__id=usuario_id)
        
        return queryset.order_by('fecha', 'hora')
